import functools
import hashlib
import hmac
import logging
import re
import time
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.hashers import check_password, make_password
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.utils.html import escape
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_POST
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError, PyMongoError

from . import email_notifications as mail
from .db import clean, db, oid
from .forms import (CandidateForm, DepartmentForm, EmployeeForm, ExemptionForm,
                    InterviewForm, LeaveForm, PTAForm, RegisterForm)
from .policy import (EXEMPTION_LIMIT, EXEMPTION_WINDOW_DAYS, LEAVE_TYPES, PENSION_RATE,
                     TAX_RATE, WORKING_DAYS_PER_MONTH)


log = logging.getLogger("hrms.security")

LATE_AFTER = "09:15"
MONTH_RE = re.compile(r"\d{4}-(0[1-9]|1[0-2])")
CANDIDATE_STAGES = ("Applied", "Interview", "Hired", "Rejected")

# ---------- roles ----------
ADMINISTRATOR, EMPLOYEE, SUPERVISOR, MANAGER, ACCOUNTANT = (
    "ADMINISTRATOR", "EMPLOYEE", "SUPERVISOR", "MANAGER", "ACCOUNTANT")
ROLES = (ADMINISTRATOR, EMPLOYEE, SUPERVISOR, MANAGER, ACCOUNTANT)

DEFAULT_ROLE = ADMINISTRATOR

LEGACY_ROLES = {"admin": ADMINISTRATOR, "employee": EMPLOYEE}   # values from the old version
ADMIN_VALUES = [ADMINISTRATOR, "admin"]                          # how admins may be stored in the DB
REVIEW_ROLES = {ADMINISTRATOR, MANAGER}
PAYROLL_ROLES = {ADMINISTRATOR, MANAGER, ACCOUNTANT}             # may see salaries / payroll totals

# ---------- security tuning ----------
MAX_LOGIN_FAILS = 5          # per email, per window
IP_LOGIN_FAILS = 25          # per IP, per window
LOGIN_WINDOW = 15 * 60
IDLE_SECONDS = getattr(settings, "SESSION_IDLE_SECONDS", 30 * 60)
ABSOLUTE_SECONDS = getattr(settings, "SESSION_ABSOLUTE_SECONDS", 8 * 3600)
MAX_PASSWORD_LEN = 128
WEAK_FRAGMENTS = ("password", "123456", "qwerty", "letmein", "welcome", "admin", "iloveyou")
NAME_RE = re.compile(r"[^<>{}$\\]{2,100}")
LEAVE_STATUSES = ("Pending", "Approved", "Rejected")

# Hash used to keep login timing identical whether or not the email exists.
DUMMY_HASH = make_password("timing-equaliser-not-a-real-password")

DB_DOWN_MSG = ("Something went wrong on our side and your request was not completed. "
               "Nothing was changed. Please try again in a moment.")


# ---------- error pages & user feedback ----------
def _error_page(request, status, title, message):
    """Friendly error page. Uses templates/error.html if you create it, otherwise a built-in page."""
    try:
        return render(request, "error.html", {"status": status, "title": title, "message": message}, status=status)
    except Exception:
        html = (f"<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
                f"<title>{escape(title)}</title>"
                f"<body style='font-family:system-ui,sans-serif;max-width:32rem;margin:15vh auto;padding:0 1rem;color:#222'>"
                f"<h1 style='font-size:1.4rem'>{status} &middot; {escape(title)}</h1>"
                f"<p>{escape(message)}</p><p><a href='/'>Back to home</a></p></body>")
        return HttpResponse(html, status=status)


def error_400(request, exception=None):
    return _error_page(request, 400, "Bad request", "We couldn't understand that request. Please go back and try again.")


def error_403(request, exception=None):
    return _error_page(request, 403, "Access denied", "You don't have permission to view this page.")


def error_404(request, exception=None):
    return _error_page(request, 404, "Page not found", "The page or record you're looking for doesn't exist or was removed.")


def error_500(request):
    return _error_page(request, 500, "Something went wrong", "An unexpected error occurred. It has been logged. Please try again shortly.")


def csrf_failed(request, reason=""):
    return _error_page(request, 403, "Form expired",
                       "Your form or session expired for security reasons. Go back, refresh the page and try again.")


def flash_form_errors(request, form):
    """Show every validation problem as a visible message (works even if the template doesn't print field errors)."""
    for field, errs in form.errors.items():
        fld = form.fields.get(field)
        label = f"{fld.label or field.replace('_', ' ').title()}: " if fld and field != "__all__" else ""
        for e in errs:
            messages.error(request, f"{label}{e}")


def form_feedback(request, form):
    """Returns the HTTP status for a form page: 400 when a submission had errors (and flashes them), else 200."""
    if request.method == "POST" and form.errors:
        flash_form_errors(request, form)
        return 400
    return 200


def db_safe(view):
    """Any MongoDB failure becomes a clear message instead of a crash page."""
    @functools.wraps(view)
    def wrapper(request, *a, **k):
        try:
            return view(request, *a, **k)
        except PyMongoError:
            log.exception("database error in %s", view.__name__)
            if request.method == "POST":
                messages.error(request, DB_DOWN_MSG)
                return safe_redirect(request, request.META.get("HTTP_REFERER"),
                                     "dashboard" if request.session.get("uid") else "landing")
            return _error_page(request, 503, "Service unavailable",
                               "We can't reach our database right now. Please try again in a few minutes.")
    return wrapper


# ---------- generic helpers ----------
def normalize_role(value):
    v = str(value or "").strip()
    if v.lower() in LEGACY_ROLES:
        return LEGACY_ROLES[v.lower()]
    return v.upper() if v.upper() in ROLES else None


def client_ip(request):
    # REMOTE_ADDR only: X-Forwarded-For is attacker-controlled unless you configure a trusted proxy.
    return request.META.get("REMOTE_ADDR", "0.0.0.0")


def clean_email(raw):
    v = raw.strip().lower() if isinstance(raw, str) else ""
    if not v or len(v) > 254:
        return ""
    try:
        validate_email(v)
    except ValidationError:
        return ""
    return v


def get_oid(pk):
    try:
        return oid(pk)
    except Exception:
        raise Http404


def safe_redirect(request, target, default):
    if target and url_has_allowed_host_and_scheme(
            target, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return redirect(target)
    return redirect(default)


def password_problems(pw, email="", name=""):
    out = []
    if len(pw) < 8:
        out.append("Password must be at least 8 characters.")
    if len(pw) > MAX_PASSWORD_LEN:
        out.append(f"Password must be at most {MAX_PASSWORD_LEN} characters.")
    if not re.search(r"[a-z]", pw):
        out.append("Add a lowercase letter.")
    if not re.search(r"[A-Z]", pw):
        out.append("Add an uppercase letter.")
    if not re.search(r"\d", pw):
        out.append("Add a number.")
    if not re.search(r"[^\w\s]", pw):
        out.append("Add a symbol (e.g. ! @ # $).")
    low = pw.lower()
    local = email.split("@")[0].lower()
    if len(local) >= 4 and local in low:
        out.append("Password must not contain your email name.")
    if any(w in low for w in WEAK_FRAGMENTS):
        out.append("Password contains a very common word or pattern.")
    if name and any(len(p) >= 4 and p.lower() in low for p in name.split()):
        out.append("Password must not contain your name.")
    try:
        validate_password(pw)          # Django's validators (common-password list etc.)
    except ValidationError as e:
        out.extend(m for m in e.messages if m not in out)
    return out


# ---------- rate limiting (stored in MongoDB, auto-expires via TTL index) ----------
def _key(scope, ident):
    return f"{scope}:{hashlib.sha256(str(ident).encode()).hexdigest()}"


def rl_hit(scope, ident, limit, window):
    """Record one event. Returns True when the caller is now over the limit."""
    key, now, col = _key(scope, ident), time.time(), db().throttle
    r = col.find_one_and_update({"_id": key, "reset": {"$gt": now}}, {"$inc": {"n": 1}},
                                return_document=ReturnDocument.AFTER)
    if r is None:
        col.replace_one({"_id": key}, {"_id": key, "n": 1, "reset": now + window,
                                       "expires": datetime.fromtimestamp(now + window, tz=timezone.utc)},
                        upsert=True)
        return 1 > limit
    return r["n"] > limit


def rl_blocked(scope, ident, limit):
    r = db().throttle.find_one({"_id": _key(scope, ident), "reset": {"$gt": time.time()}})
    return bool(r and r["n"] >= limit)


def rl_clear(scope, ident):
    db().throttle.delete_one({"_id": _key(scope, ident)})


# ---------- audit log ----------
def audit(request, action, target="", actor=None):
    try:
        db().audit_log.insert_one({
            "time": datetime.now(timezone.utc),
            "actor": actor or request.session.get("email", ""),
            "role": request.session.get("role", ""),
            "ip": client_ip(request), "action": action, "target": str(target)[:200]})
    except Exception:
        log.exception("audit write failed")


# ---------- sessions & access control ----------
def terminated(email):
    e = db().employees.find_one({"email": email}, {"status": 1})
    return bool(e and e.get("status") == "Terminated")


def current_user(request):
    """Validate the session against the database on EVERY request.
    Role changes, password changes, deleted or terminated users take effect immediately."""
    uid = request.session.get("uid")
    if not uid:
        return None
    now = int(time.time())
    if (now - request.session.get("last_seen", now) > IDLE_SECONDS or
            now - request.session.get("login_at", now) > ABSOLUTE_SECONDS):
        request.session.flush()
        messages.info(request, "Your session expired. Please sign in again.")
        return None
    try:
        u = db().users.find_one({"_id": oid(uid)})
    except PyMongoError:
        raise                                                   # DB trouble is handled by db_safe, not a forced logout
    except Exception:
        u = None
    if (not u or u.get("pwv", 0) != request.session.get("pwv", 0) or terminated(u["email"])):
        request.session.flush()
        messages.warning(request, "You've been signed out. Please sign in again.")
        return None
    u["role"] = normalize_role(u.get("role")) or EMPLOYEE      # unknown value => least privilege
    request.session.update({"last_seen": now, "role": u["role"], "name": u["name"], "email": u["email"]})
    return u


def login_required(view):
    cached = never_cache(view)

    @functools.wraps(view)
    def wrapper(request, *a, **k):
        u = current_user(request)
        if not u:
            return redirect("login")
        request.hr_user, request.role = u, u["role"]
        return cached(request, *a, **k)
    return db_safe(wrapper)


def role_required(*allowed):
    def deco(view):
        @login_required
        @functools.wraps(view)
        def wrapper(request, *a, **k):
            if request.role not in allowed:
                audit(request, "access_denied", request.path)
                messages.error(request, "You don't have permission to do that. Ask an administrator if you need access.")
                return redirect("dashboard")
            return view(request, *a, **k)
        return wrapper
    return deco


admin_required = role_required(ADMINISTRATOR)
review_required = role_required(*REVIEW_ROLES)


def is_admin(request):
    return getattr(request, "role", None) == ADMINISTRATOR


def can_review(request):
    return getattr(request, "role", None) in REVIEW_ROLES


def can_see_pay(request):
    return getattr(request, "role", None) in PAYROLL_ROLES


def page(request, template, nav, status=200, filter_status=None, **ctx):
    if filter_status is not None:
        ctx["status"] = filter_status          # template variable, not the HTTP code
    ctx.update(nav=nav, currency=settings.CURRENCY, role=getattr(request, "role", None),
               is_admin=is_admin(request), can_review=can_review(request), can_see_pay=can_see_pay(request),
               can_payroll=getattr(request, "role", None) in (ADMINISTRATOR, ACCOUNTANT))
    return render(request, template, ctx, status=status)


def count_admins():
    return db().users.count_documents({"role": {"$in": ADMIN_VALUES}})


# ---------- department-manager scoping ----------
# An administrator assigns one manager per department (departments.manager_email).
# A MANAGER only sees and reviews employees, attendance and leave of the departments assigned to them.
def managed_departments(request):
    """None = unrestricted (administrator). Otherwise the names of the departments this manager runs."""
    if is_admin(request):
        return None
    if request.role == MANAGER:
        return [d["name"] for d in db().departments.find({"manager_email": request.session["email"]}, {"name": 1})]
    return []


def team_emails(request):
    """None = everyone (administrator). Otherwise the e-mails of employees in the departments this manager runs."""
    deps = managed_departments(request)
    if deps is None:
        return None
    return {e["email"] for e in db().employees.find({"department": {"$in": deps}}, {"email": 1})}


def scoped(request, include_self=True):
    me = request.session["email"]
    if is_admin(request):
        return {}
    if request.role == MANAGER:
        team = team_emails(request)
    elif request.role == SUPERVISOR:
        team = {e["email"] for e in db().employees.find({"supervisor_email": me}, {"email": 1})}
    else:
        return {"email": me}
    team.add(me) if include_self else team.discard(me)
    return {"email": {"$in": sorted(team)}}


def can_manage(request, email):
    team = team_emails(request)
    return team is None or email in team


def manager_candidates():
    """Active employees an administrator may assign as department manager."""
    return list(db().employees.find({"status": "Active", "role": MANAGER}).sort("first_name", 1))


def departments_run_by(email):
    return [d["name"] for d in db().departments.find({"manager_email": email}, {"name": 1})]


# ---------- approval workflow: employee -> supervisor -> HR ----------
def supervisor_of(email):
    e = db().employees.find_one({"email": email}, {"supervisor_email": 1}) or {}
    sup = e.get("supervisor_email", "")
    if sup and sup != email and db().employees.find_one({"email": sup, "status": "Active", "role": SUPERVISOR}):
        return sup
    return ""                                                   # no supervisor -> goes straight to HR


def new_request(request, **fields):
    me, sup = request.session["email"], supervisor_of(request.session["email"])
    return {"email": me, "name": request.session["name"], "status": "Pending",
            "stage": "supervisor" if sup else "hr", "supervisor_email": sup,
            "created": datetime.now().isoformat(timespec="seconds"), **fields}


def can_act(request, doc):
    """Is this request waiting for the current user? HR (administrator) can act on any pending request."""
    me = request.session["email"]
    if doc.get("status") != "Pending" or doc["email"] == me:
        return False
    if is_admin(request):
        return True
    if doc.get("stage", "hr") == "supervisor":
        return request.role == SUPERVISOR and doc.get("supervisor_email") == me
    return False

def awaiting(request):
    """Mongo filter for requests waiting on this user (None if they approve nothing)."""
    me = request.session["email"]
    if request.role == SUPERVISOR:
        return {"status": "Pending", "stage": "supervisor", "supervisor_email": me, "email": {"$ne": me}}
    if is_admin(request):
        return {"status": "Pending", "stage": {"$in": ["hr", None]}, "email": {"$ne": me}}
    return None



def decide(request, col, pk, action, label, on_final=None, back="dashboard"):
    status = {"approve": "Approved", "reject": "Rejected"}.get(action)
    if not status:
        raise Http404
    c = db()[col]
    doc = c.find_one({"_id": get_oid(pk)})
    if not doc:
        messages.error(request, f"That {label} no longer exists.")
    elif doc["status"] != "Pending":
        messages.info(request, f"That {label} was already {doc['status'].lower()}.")
    elif doc["email"] == request.session["email"]:
        messages.error(request, f"You can't review your own {label}.")
    elif not can_act(request, doc):
        audit(request, "access_denied", f"{col} of {doc['email']}")
        messages.error(request, "That request isn't waiting for your approval.")
    else:
        hr = is_admin(request)
        stage = "hr" if hr else doc.get("stage", "hr")          # who is acting right now
        now = datetime.now().isoformat(timespec="seconds")
        step = {f"{stage}_by": request.session["name"], f"{stage}_at": now}
        if status == "Rejected":
            upd = {"status": "Rejected", "stage": "done", **step}
        elif stage == "supervisor" and doc.get("needs_hr", True):
            upd = {"stage": "hr", **step}                       # supervisor ok -> now HR
        else:
            upd = {"status": "Approved", "stage": "done", **step}   # HR approval is always final
        flt = {"_id": doc["_id"], "status": "Pending"}          # atomic: no double decisions
        if "stage" in doc:
            flt["stage"] = doc["stage"]
        if c.update_one(flt, {"$set": upd}).modified_count:
            audit(request, f"{col}_{action}_{stage}", f"{doc['email']}")
            if upd.get("status") == "Approved" and on_final:
                on_final({**doc, **upd})
            if upd.get("stage") == "hr":
                messages.success(request, "Approved. It is now waiting for HR.")
            else:
                messages.success(request, f"{doc['name']}'s {label} was {status.lower()}.")
        else:
            messages.info(request, "Someone else already decided that request.")
    return safe_redirect(request, request.POST.get("next"), back)

def weekdays(start, end):
    d = start
    while d <= end:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def month_str(d=None):
    return (d or date.today()).strftime("%Y-%m")


def month_bounds(m):
    y, mo = map(int, m.split("-"))
    return date(y, mo, 1), date(y, mo, monthrange(y, mo)[1])


def parse_day(raw, default):
    try:
        return date.fromisoformat(raw).isoformat()
    except (TypeError, ValueError):
        return default


def xlsx_response(filename, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":                           # stop user text like "=1+1" becoming a formula
                cell.data_type = "s"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(40, max(len(str(c.value or "")) for c in col) + 2)
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(resp)
    return resp


def leave_to_attendance(l):
    """HR approved: mark every working day of the leave as present ('On leave'). Real clock records are kept."""
    for d in weekdays(date.fromisoformat(l["start_date"]), date.fromisoformat(l["end_date"])):
        db().attendance.update_one(
            {"email": l["email"], "date": d.isoformat()},
            {"$setOnInsert": {"name": l["name"], "check_in": "", "check_out": "", "status": "On leave",
                              "leave_type": l["leave_type"], "leave_id": str(l["_id"])}},
            upsert=True)


def department_manager(dep_name):
    """(department, manager employee) for a department name; either may be None."""
    dept = db().departments.find_one({"name": dep_name}) if dep_name else None
    mgr = None
    if dept and dept.get("manager_email"):
        mgr = db().employees.find_one({"email": dept["manager_email"], "status": "Active"})
    return dept, mgr


_seeded = False


def seed():
    global _seeded
    if _seeded:
        return
    d = db()
    d.users.create_index("email", unique=True)
    d.employees.create_index("email", unique=True)
    d.departments.create_index("name", unique=True)
    d.throttle.create_index("expires", expireAfterSeconds=0)
    d.audit_log.create_index("time")
    d.payroll.create_index([("month", 1), ("email", 1)], unique=True)
    try:
        d.attendance.create_index([("email", 1), ("date", 1)], unique=True)
    except Exception:
        log.exception("attendance unique index not created (duplicate rows exist?)")
    if not d.departments.count_documents({}):
        for n in ("Engineering", "Human Resources", "Finance", "Marketing", "Operations"):
            d.departments.insert_one({"name": n, "head": "", "manager_email": "", "description": ""})
    _seeded = True


# ---------- landing ----------
def landing(request):
    if request.session.get("uid"):
        return redirect("dashboard")
    return render(request, "landing.html")


# ---------- auth ----------
@db_safe
def login_view(request):
    if current_user(request):
        return redirect("dashboard")
    ctx = {"first": False, "email": ""}
    try:
        seed()
        ctx["first"] = db().users.count_documents({}) == 0
    except Exception:
        log.exception("MongoDB unreachable")
        messages.error(request, "The service is temporarily unavailable. Please try again shortly.")
        return render(request, "login.html", ctx, status=503)

    if request.method == "POST":
        email = ctx["email"] = clean_email(request.POST.get("email"))
        password = request.POST.get("password", "")
        if not isinstance(password, str):
            password = ""
        ip = client_ip(request)

        if not email:
            messages.error(request, "Enter a valid email address.")
            return render(request, "login.html", ctx, status=400)
        if not password:
            messages.error(request, "Enter your password.")
            return render(request, "login.html", ctx, status=400)

        if rl_blocked("login-ip", ip, IP_LOGIN_FAILS) or rl_blocked("login-email", email, MAX_LOGIN_FAILS):
            audit(request, "login_blocked", email, actor=email)
            messages.error(request, "Too many failed attempts. Please wait 15 minutes and try again.")
            return render(request, "login.html", ctx, status=429)

        u = db().users.find_one({"email": email})
        ok = check_password(password[:MAX_PASSWORD_LEN], u["password"] if u else DUMMY_HASH)

        if u and ok and not terminated(u["email"]):
            rl_clear("login-email", email)
            request.session.cycle_key()                       # new session id on login (anti-fixation)
            now = int(time.time())
            request.session.update({
                "uid": str(u["_id"]), "name": u["name"], "email": u["email"],
                "role": normalize_role(u.get("role")) or EMPLOYEE,
                "pwv": u.get("pwv", 0), "login_at": now, "last_seen": now})
            audit(request, "login_ok", email, actor=email)
            messages.success(request, f"Welcome back, {u['name']}.")
            return redirect("dashboard")

        rl_hit("login-ip", ip, IP_LOGIN_FAILS, LOGIN_WINDOW)
        rl_hit("login-email", email, MAX_LOGIN_FAILS, LOGIN_WINDOW)
        audit(request, "login_fail", email, actor=email)
        messages.error(request, "Email or password is incorrect.")      # same message for every failure
        return render(request, "login.html", ctx, status=401)
    return render(request, "login.html", ctx)


@db_safe
def register_view(request):
    """The first account becomes ADMINISTRATOR. After that only employees HR has added can register.
    Every new account gets DEFAULT_ROLE (ADMINISTRATOR) unless HR set a role on the employee record."""
    if current_user(request):
        return redirect("dashboard")
    try:
        seed()
        first = db().users.count_documents({}) == 0
    except Exception:
        log.exception("MongoDB unreachable")
        messages.error(request, "The service is temporarily unavailable. Please try again shortly.")
        return render(request, "register.html", {"form": RegisterForm(first=False), "first": False}, status=503)

    form = RegisterForm(request.POST or None, first=first)
    if request.method == "POST":
        if rl_hit("register-ip", client_ip(request), 10, 3600):
            messages.error(request, "Too many registration attempts. Please try again in an hour.")
            return render(request, "register.html", {"form": form, "first": first}, status=429)

        if form.is_valid():
            c = form.cleaned_data
            email = clean_email(c.get("email"))
            name, role = (c.get("name") or "").strip(), DEFAULT_ROLE
            if not email:
                form.add_error("email", "Enter a valid email address.")

            code = getattr(settings, "REGISTRATION_INVITE_CODE", "")
            if code and not hmac.compare_digest(request.POST.get("invite_code", "").encode(), code.encode()):
                form.add_error(None, "A valid invitation code is required.")

            if email and not first:
                emp = db().employees.find_one({"email": email})
                if not emp or emp.get("status") == "Terminated":
                    form.add_error("email", "Registration isn't available for this email. Ask HR to add you first.")
                else:
                    name = f"{emp['first_name']} {emp['last_name']}"
                    role = normalize_role(emp.get("role")) or DEFAULT_ROLE
            elif first and not NAME_RE.fullmatch(name):
                form.add_error("name", "Enter your full name (2-100 characters, no special symbols).")

            for p in password_problems(c.get("password", ""), email, name):
                form.add_error("password", p)

            if not form.errors:
                claimed = False
                try:
                    if first:                                   # atomic claim: only one first-admin can ever win
                        try:
                            db().meta.insert_one({"_id": "bootstrap"})
                            claimed = True
                        except DuplicateKeyError:
                            # A lock exists but there are no users: it's an orphan from a failed
                            # earlier attempt (or a manually emptied users collection). Reclaim it.
                            if db().users.count_documents({}) == 0:
                                db().meta.delete_one({"_id": "bootstrap"})
                                db().meta.insert_one({"_id": "bootstrap"})
                                claimed = True
                            else:
                                raise
                    db().users.insert_one({
                        "email": email, "name": name, "role": role, "pwv": 0,
                        "password": make_password(c["password"]),
                        "created": datetime.now(timezone.utc).isoformat(timespec="seconds")})
                except DuplicateKeyError:
                    if claimed:
                        db().meta.delete_one({"_id": "bootstrap"})
                    form.add_error(None, "An account with this email already exists. Please sign in.")
                except PyMongoError:
                    if claimed:
                        db().meta.delete_one({"_id": "bootstrap"})
                    raise                                       # db_safe shows the friendly message
                else:
                    audit(request, "register", f"{email} as {role}", actor=email)
                    mail.account_created(email, name, role)     # confirmation e-mail (background, never blocks)
                    messages.success(request, f"Account created for {name} ({role.title()}). You can now sign in with {email}.")
                    return redirect("login")
        return render(request, "register.html", {"form": form, "first": first}, status=form_feedback(request, form))
    return render(request, "register.html", {"form": form, "first": first})


def logout_view(request):
    audit(request, "logout")
    request.session.flush()
    messages.success(request, "You've been signed out.")
    return redirect("landing")


@login_required
def password_change(request):
    status = 200
    if request.method == "POST":
        u, uid = request.hr_user, request.session["uid"]
        current, new = request.POST.get("current", ""), request.POST.get("new", "")
        confirm = request.POST.get("confirm")
        if rl_blocked("pwchange", uid, 5):
            messages.error(request, "Too many failed attempts. Try again in 15 minutes.")
            status = 429
        elif not current or not new:
            messages.error(request, "Enter your current password and a new password.")
            status = 400
        elif not check_password(current[:MAX_PASSWORD_LEN], u["password"]):
            rl_hit("pwchange", uid, 5, LOGIN_WINDOW)
            audit(request, "password_change_fail")
            messages.error(request, "Current password is incorrect.")
            status = 400
        elif confirm is not None and confirm != new:
            messages.error(request, "New passwords do not match.")
            status = 400
        elif check_password(new[:MAX_PASSWORD_LEN], u["password"]):
            messages.error(request, "New password must be different from the current one.")
            status = 400
        elif (problems := password_problems(new, u["email"], u["name"])):
            for p in problems:
                messages.error(request, p)
            status = 400
        else:
            pwv = u.get("pwv", 0) + 1                           # invalidates every other session of this user
            db().users.update_one({"_id": u["_id"]}, {"$set": {"password": make_password(new), "pwv": pwv}})
            request.session["pwv"] = pwv
            request.session.cycle_key()
            rl_clear("pwchange", uid)
            audit(request, "password_changed")
            messages.success(request, "Password updated. Other devices have been signed out.")
            return redirect("dashboard")
    return page(request, "password.html", "", status=status)


# ---------- dashboard ----------
@login_required
def dashboard(request):
    d, today = db(), date.today().isoformat()
    me = request.session["email"]
    deps = managed_departments(request) if request.role == MANAGER else None
    emp_scope = {} if deps is None else {"department": {"$in": deps}}
    groups = list(d.employees.aggregate([
        {"$match": {"status": "Active", **emp_scope}},
        {"$group": {"_id": "$department", "n": {"$sum": 1}}}, {"$sort": {"n": -1}}]))
    top = max([g["n"] for g in groups] or [1])
    active = list(d.employees.find({"status": "Active", **emp_scope}, {"salary": 1}))
    reviewer = can_review(request)
    proj = None if can_see_pay(request) else {"salary": 0}
    team = scoped(request, include_self=False)
    return page(request, "dashboard.html", "dashboard",
                headcount=len(active),
                payroll=sum(e.get("salary", 0) for e in active) if can_see_pay(request) else None,
                departments=d.departments.count_documents({} if deps is None else {"name": {"$in": deps}}),
                pending=d.leaves.count_documents({"status": "Pending", **team}) if reviewer else 0,
                present=d.attendance.count_documents({"date": today, **scoped(request)}),
                my_pending=d.leaves.count_documents({"email": me, "status": "Pending"}),
                roster=[{"name": g["_id"], "n": g["n"], "pct": max(6, int(g["n"] * 100 / top))} for g in groups],
                recent=[clean(e) for e in d.employees.find(emp_scope, proj).sort("hire_date", -1).limit(5)],
                waiting=[clean(l) for l in d.leaves.find({"status": "Pending", **team}).sort("created", 1).limit(5)]
                if reviewer else [])


# ---------- employees ----------
@login_required
def employees(request):
    q, dep = request.GET.get("q", "").strip()[:100], request.GET.get("department", "").strip()[:100]
    f = {}
    if q:
        f["$or"] = [{k: {"$regex": re.escape(q), "$options": "i"}} for k in ("first_name", "last_name", "email", "position")]
    deps = managed_departments(request) if request.role == MANAGER else None
    if deps is not None:                                        # a manager only sees their own departments
        f["department"] = {"$in": ([dep] if dep in deps else []) if dep else deps}
    elif dep:
        f["department"] = dep
    proj = None if can_see_pay(request) else {"salary": 0}
    rows = [clean(e) for e in db().employees.find(f, proj).sort("first_name", 1)]
    if (q or dep) and not rows:
        messages.info(request, "No employees match your search. Try different keywords or clear the filters.")
    return page(request, "employees.html", "employees", rows=rows, q=q, dep=dep,
                departments=sorted(deps) if deps is not None else [x["name"] for x in db().departments.find().sort("name", 1)])


@admin_required
def employee_edit(request, pk=None):
    doc = None
    if pk:
        doc = db().employees.find_one({"_id": get_oid(pk)})
        if not doc:
            raise Http404
    form = EmployeeForm(request.POST or None, initial=doc)

    form.fields["supervisor_email"].choices = [("", "No supervisor")] + [
        (s["email"], f"{s['first_name']} {s['last_name']}")
        for s in db().employees.find({"status": "Active", "role": SUPERVISOR}).sort("first_name", 1)]

    if request.method == "POST" and form.is_valid():
        data = dict(form.cleaned_data)
        for k, v in data.items():
            if isinstance(v, str):
                data[k] = v.strip()
        email = clean_email(data.get("email"))
        salary = float(data["salary"])
        hire = data["hire_date"]
        role = normalize_role(data.get("role")) if data.get("role") else None

        if not email:
            form.add_error("email", "Enter a valid email address.")
        if not 0 <= salary <= 1_000_000_000:
            form.add_error("salary", "Enter a realistic salary.")
        if not date(1950, 1, 1) <= hire <= date.today() + timedelta(days=366):
            form.add_error("hire_date", "Hire date is out of range.")
        if "role" in form.fields and data.get("role") and not role:
            form.add_error("role", "Choose a valid role.")
        if doc and email and email != doc["email"] and db().users.find_one({"email": email}):
            form.add_error("email", "That email already has a login account.")
        if doc and role and role != ADMINISTRATOR:
            u = db().users.find_one({"email": doc["email"]}, {"role": 1})
            if u and u.get("role") in ADMIN_VALUES and count_admins() <= 1:
                form.add_error("role", "You can't demote the last administrator.")
        if doc:                                                 # don't orphan a department that this person manages
            runs = departments_run_by(doc["email"])
            new_role = role or normalize_role(doc.get("role"))
            if runs and (new_role != MANAGER or data.get("status") == "Terminated" or email != doc["email"]):
                form.add_error(None, f"{doc['first_name']} manages {', '.join(runs)}. "
                                     f"Assign another manager to that department first.")

        if not form.errors:
            data.update(hire_date=hire.isoformat(), salary=salary, email=email)
            if role:
                data["role"] = role
            fullname = f"{data['first_name']} {data['last_name']}"
            try:
                if doc:
                    db().employees.update_one({"_id": doc["_id"]}, {"$set": data})
                    upd = {"email": email, "name": fullname}
                    if role:
                        upd["role"] = role
                    db().users.update_one({"email": doc["email"]}, {"$set": upd})
                else:
                    db().employees.insert_one(data)
            except DuplicateKeyError:
                form.add_error("email", "Someone with this email already exists.")
            else:
                audit(request, "employee_saved" if doc else "employee_added", email)
                dept, mgr = department_manager(data.get("department"))
                if not doc:                                     # e-mail the new employee (welcome + how to register)
                    mail.employee_welcome(data, mgr)
                if (not doc or doc.get("department") != data.get("department")) and dept:
                    if mgr and mgr["email"] != email:           # e-mail the department manager
                        mail.manager_new_member(mgr, data, dept["name"])
                    elif not mgr:
                        messages.warning(request, f"{dept['name']} has no manager yet, so only administrators can "
                                                  f"review {fullname}'s leave and attendance. Assign one under Departments.")
                if doc:
                    messages.success(request, f"{fullname}'s details were updated.")
                else:
                    messages.success(request, f"{fullname} was added to {data.get('department') or 'the company'}. "
                                              f"A welcome e-mail with registration instructions was sent to {email}.")
                return redirect("employees")
    return page(request, "form.html", "employees", status=form_feedback(request, form), form=form, back="employees",
                title="Edit employee" if doc else "Add employee", save="Save employee")


@admin_required
@require_POST
def employee_reset(request, pk):
    e = db().employees.find_one({"_id": get_oid(pk)})
    if not e:
        messages.error(request, "That employee no longer exists.")
        return redirect("employees")
    name = f"{e['first_name']} {e['last_name']}"
    if e["email"] == request.session["email"]:
        messages.error(request, "Use 'Change password' to reset your own login.")
    elif db().users.delete_one({"email": e["email"], "role": {"$nin": ADMIN_VALUES}}).deleted_count:
        audit(request, "login_reset", e["email"])
        messages.success(request, f"{name}'s login was reset. They can register again with their work email.")
    else:
        messages.info(request, f"{name} hasn't registered yet, or is an administrator (not resettable here).")
    return redirect("employees")


@admin_required
@require_POST
def employee_delete(request, pk):
    e = db().employees.find_one({"_id": get_oid(pk)})
    if not e:
        messages.error(request, "That employee no longer exists.")
        return redirect("employees")
    u = db().users.find_one({"email": e["email"]}, {"role": 1})
    runs = departments_run_by(e["email"])
    if e["email"] == request.session["email"]:
        messages.error(request, "You can't remove your own account.")
    elif u and u.get("role") in ADMIN_VALUES and count_admins() <= 1:
        messages.error(request, "You can't remove the last administrator.")
    elif runs:
        messages.error(request, f"{e['first_name']} manages {', '.join(runs)}. Assign another manager first.")
    else:
        db().employees.delete_one({"_id": e["_id"]})
        db().users.delete_one({"email": e["email"]})
        audit(request, "employee_deleted", e["email"])
        messages.success(request, f"{e['first_name']} {e['last_name']} was removed.")
    return redirect("employees")


# ---------- departments ----------
@login_required
def departments(request):
    deps = managed_departments(request) if request.role == MANAGER else None
    f = {} if deps is None else {"name": {"$in": deps}}
    rows = []
    for dp in db().departments.find(f).sort("name", 1):
        dp["count"] = db().employees.count_documents({"department": dp["name"], "status": "Active"})
        rows.append(clean(dp))
    return page(request, "departments.html", "departments", rows=rows)


@admin_required
def department_edit(request, pk=None):
    doc = None
    if pk:
        doc = db().departments.find_one({"_id": get_oid(pk)})
        if not doc:
            raise Http404
    form = DepartmentForm(request.POST or None, initial=doc)
    if "manager_email" in form.fields:                          # the administrator picks the manager from active MANAGER employees
        form.fields["manager_email"].choices = [("", "No manager assigned")] + [
            (m["email"], f"{m['first_name']} {m['last_name']} ({m['email']})") for m in manager_candidates()]
    if request.method == "POST" and form.is_valid():
        data = {k: (v.strip() if isinstance(v, str) else v) for k, v in form.cleaned_data.items()}
        if not NAME_RE.fullmatch(data.get("name", "")):
            form.add_error("name", "Enter a name of 2-100 characters without special symbols.")
        if len(str(data.get("description", ""))) > 500:
            form.add_error("description", "Keep the description under 500 characters.")
        mgr = None
        if "manager_email" in form.fields:
            mgr_email = clean_email(data.get("manager_email")) if data.get("manager_email") else ""
            if mgr_email:
                mgr = db().employees.find_one({"email": mgr_email, "status": "Active", "role": MANAGER})
                if not mgr:
                    form.add_error("manager_email", "Choose an active employee whose role is Manager.")
            data["manager_email"] = mgr_email
            if not form.errors:
                data["head"] = f"{mgr['first_name']} {mgr['last_name']}" if mgr else ""
        if not form.errors:
            try:
                if doc:
                    db().departments.update_one({"_id": doc["_id"]}, {"$set": data})
                    if doc["name"] != data["name"]:
                        db().employees.update_many({"department": doc["name"]}, {"$set": {"department": data["name"]}})
                else:
                    db().departments.insert_one(data)
            except DuplicateKeyError:
                form.add_error("name", "A department with this name already exists.")
            else:
                audit(request, "department_saved", data["name"])
                if mgr and mgr["email"] != (doc or {}).get("manager_email", ""):
                    audit(request, "department_manager_assigned", f"{data['name']} -> {mgr['email']}")
                    mail.manager_assigned(mgr, data["name"],
                                          db().employees.count_documents({"department": data["name"], "status": "Active"}))
                messages.success(request, f"Department \"{data['name']}\" was " + ("updated." if doc else "created.")
                                 + (f" {mgr['first_name']} {mgr['last_name']} is now its manager." if mgr else ""))
                return redirect("departments")
    return page(request, "form.html", "departments", status=form_feedback(request, form), form=form, back="departments",
                title="Edit department" if doc else "Add department", save="Save department")


@admin_required
@require_POST
def department_delete(request, pk):
    dp = db().departments.find_one({"_id": get_oid(pk)})
    if not dp:
        messages.error(request, "That department no longer exists.")
    elif db().employees.count_documents({"department": dp["name"]}):
        messages.error(request, f"\"{dp['name']}\" still has employees. Move or remove them first.")
    else:
        db().departments.delete_one({"_id": dp["_id"]})
        audit(request, "department_deleted", dp["name"])
        messages.success(request, f"Department \"{dp['name']}\" was deleted.")
    return redirect("departments")


# ---------- leave ----------
@login_required
def leaves(request):
    status = request.GET.get("status", "")
    status = status if status in LEAVE_STATUSES else ""
    f = scoped(request)
    if status:
        f["status"] = status
    rows = []
    for l in db().leaves.find(f).sort("created", -1).limit(500):
        l = clean(l)
        l["can_act"] = can_act(request, l)
        rows.append(l)
    return page(request, "leaves.html", "leaves", rows=rows, filter_status=status,
                tabs=["", *LEAVE_STATUSES])


@login_required
def leave_apply(request):
    form = LeaveForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        c, me = form.cleaned_data, request.session["email"]
        start, end, reason = c["start_date"], c["end_date"], (c.get("reason") or "").strip()
        days = (end - start).days + 1
        if end < start:
            form.add_error("end_date", "End date can't be before the start date.")
        elif days > 365:
            form.add_error("end_date", "A single request can't exceed 365 days.")
       
        if start > date.today() + timedelta(days=730):
            form.add_error("start_date", "Start date is too far in the future.")
        if len(reason) > 1000:
            form.add_error("reason", "Keep the reason under 1000 characters.")
        if not form.errors:
            cap = LEAVE_TYPES.get(c["leave_type"])
            if cap is not None:                                 # yearly entitlement per leave type
                used = sum(x["days"] for x in db().leaves.find(
                    {"email": me, "leave_type": c["leave_type"], "status": {"$in": ["Pending", "Approved"]},
                     "start_date": {"$gte": f"{start.year}-01-01", "$lte": f"{start.year}-12-31"}}, {"days": 1}))
                if used + days > cap:
                    form.add_error("leave_type", f"{c['leave_type']} leave allows {cap} days a year; you have {max(cap - used, 0)} left.")
        if not form.errors:
            if db().leaves.count_documents({"email": me, "status": "Pending"}) >= 10:
                form.add_error(None, "You have too many pending requests. Wait for a decision first.")
            elif db().leaves.find_one({"email": me, "status": {"$in": ["Pending", "Approved"]},
                                       "start_date": {"$lte": end.isoformat()},
                                       "end_date": {"$gte": start.isoformat()}}):
                form.add_error(None, "These dates overlap another pending or approved request.")
        if not form.errors:
            db().leaves.insert_one(new_request(
                request, leave_type=c["leave_type"], start_date=start.isoformat(), end_date=end.isoformat(),
                days=days, reason=reason, needs_hr=True))
            audit(request, "leave_requested", f"{start}..{end}")
            who = "your supervisor" if supervisor_of(me) else "HR"
            messages.success(request, f"Leave request submitted: {days} day{'s' if days != 1 else ''} "
                                      f"({start:%d %b %Y} to {end:%d %b %Y}). It is now with {who}.")
            return redirect("leaves")
    return page(request, "form.html", "leaves", status=form_feedback(request, form), form=form, back="leaves",
                title="Request leave", save="Submit request")


@login_required
@require_POST
def leave_action(request, pk, action):
    return decide(request, "leaves", pk, action, "leave request", leave_to_attendance, "leaves")


# ---------- attendance ----------
@login_required
def attendance(request):
    today = date.today().isoformat()
    raw = request.GET.get("date", "")
    try:
        day = date.fromisoformat(raw).isoformat() if raw else today
    except ValueError:
        day = today
        messages.warning(request, "That date wasn't valid, so today's attendance is shown.")
    rows = []
    for a in db().attendance.find({"date": day, **scoped(request)}).sort("check_in", 1).limit(1000):
        a["late"] = bool(a.get("check_in")) and a["check_in"] > LATE_AFTER
        rows.append(clean(a))
    mine = db().attendance.find_one({"email": request.session["email"], "date": today})
    notices = [clean(n) for n in db().notices.find().sort("created", -1).limit(5)]
    return page(request, "attendance.html", "attendance", rows=rows, day=day, today=today, mine=mine,
                late_after=LATE_AFTER, notices=notices, month_start=today[:8] + "01")


@login_required
@require_POST
def clock(request):
    now, today, email = datetime.now().strftime("%H:%M"), date.today().isoformat(), request.session["email"]
    rec = db().attendance.find_one({"email": email, "date": today})
    if rec and rec.get("status") == "On leave":                 # approved leave counts as present
        messages.info(request, "You're on approved leave today, so there's nothing to clock.")
        return redirect("attendance")
    if not rec:
        try:
            db().attendance.insert_one({"email": email, "name": request.session["name"], "date": today,
                                        "check_in": now, "check_out": None})
            late = now > LATE_AFTER
            messages.success(request, f"Clocked in at {now}." + (f" You're marked late (after {LATE_AFTER})." if late else ""))
        except DuplicateKeyError:                               # double-click / race
            messages.info(request, "You're already clocked in.")
    elif not rec.get("check_out"):
        res = db().attendance.update_one({"_id": rec["_id"], "check_out": None}, {"$set": {"check_out": now}})
        if res.modified_count:
            messages.success(request, f"Clocked out at {now}. Have a good evening.")
        else:
            messages.info(request, "You've already clocked out.")
    else:
        messages.info(request, "You've already completed today's attendance.")
    return redirect("attendance")


@login_required
def attendance_export(request):
    today = date.today()
    start = parse_day(request.GET.get("from"), today.replace(day=1).isoformat())
    end = parse_day(request.GET.get("to"), today.isoformat())
    out = []
    for a in db().attendance.find({"date": {"$gte": start, "$lte": end}, **scoped(request)}) \
            .sort([("date", 1), ("name", 1)]).limit(20000):
        ci = a.get("check_in") or ""
        out.append([a["date"], a.get("name", ""), a["email"], ci, a.get("check_out") or "",
                    a.get("status") or ("Late" if ci > LATE_AFTER else "On time")])
    audit(request, "attendance_export", f"{start}..{end}")
    return xlsx_response(f"attendance_{start}_to_{end}.xlsx",
                         ["Date", "Name", "Email", "In", "Out", "Status"], out)


@admin_required
@require_POST
def notice_post(request):
    title, body = request.POST.get("title", "").strip()[:120], request.POST.get("body", "").strip()[:1000]
    if not title:
        messages.error(request, "A notice needs a title.")
    else:
        db().notices.insert_one({"title": title, "body": body, "by": request.session["name"],
                                 "created": datetime.now().isoformat(timespec="seconds")})
        audit(request, "notice_posted", title)
        messages.success(request, "Notice posted on the attendance page.")
    return redirect("attendance")


@admin_required
@require_POST
def notice_delete(request, pk):
    db().notices.delete_one({"_id": get_oid(pk)})
    audit(request, "notice_deleted", pk)
    return redirect("attendance")


# ---------- exemptions ----------
def exemptions_used(email, month):
    return db().exemptions.count_documents({"email": email, "month": month, "status": {"$in": ["Pending", "Approved"]}})


def exemption_to_attendance(x):
    fix = {k: x[k] for k in ("check_in", "check_out") if x.get(k)}
    db().attendance.update_one({"email": x["email"], "date": x["date"]},
                               {"$set": {**fix, "status": "Exempted"}, "$setOnInsert": {"name": x["name"]}},
                               upsert=True)


@login_required
def exemptions(request):
    rows = []
    for x in db().exemptions.find(scoped(request)).sort("created", -1).limit(500):
        x = clean(x)
        x["can_act"] = can_act(request, x)
        rows.append(x)
    return page(request, "exemptions.html", "exemptions", rows=rows,
                used=exemptions_used(request.session["email"], month_str()), limit=EXEMPTION_LIMIT)


@login_required
def exemption_apply(request):
    form = ExemptionForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        c, me = form.cleaned_data, request.session["email"]
        day, kind = c["date"], c["kind"]
        if day > date.today():
            form.add_error("date", "You can't apply for a future date.")
        elif day < date.today() - timedelta(days=EXEMPTION_WINDOW_DAYS):
            form.add_error("date", f"Requests must be made within {EXEMPTION_WINDOW_DAYS} days.")
        if kind in ("in", "both") and not c.get("check_in"):
            form.add_error("check_in", "Enter the time you arrived.")
        if kind in ("out", "both") and not c.get("check_out"):
            form.add_error("check_out", "Enter the time you left.")
        if not form.errors and db().exemptions.find_one(
                {"email": me, "date": day.isoformat(), "status": {"$in": ["Pending", "Approved"]}}):
            form.add_error("date", "You already have a request for that day.")
        if not form.errors:
            month = day.strftime("%Y-%m")
            over = exemptions_used(me, month) >= EXEMPTION_LIMIT
            db().exemptions.insert_one(new_request(
                request, date=day.isoformat(), month=month, kind=kind,
                kind_label=dict(form.fields["kind"].choices)[kind],
                check_in=c["check_in"].strftime("%H:%M") if c.get("check_in") else "",
                check_out=c["check_out"].strftime("%H:%M") if c.get("check_out") else "",
                reason=c["reason"].strip(), over_limit=over, needs_hr=over))
            audit(request, "exemption_requested", f"{day} ({kind})")
            if over:
                messages.warning(request, f"You've passed the {EXEMPTION_LIMIT} exemptions allowed for {month}. "
                                          f"This one needs supervisor and HR approval.")
            else:
                messages.success(request, "Exemption request submitted.")
            return redirect("exemptions")
    return page(request, "form.html", "exemptions", status=form_feedback(request, form), form=form,
                back="exemptions", title="Missed clock-in / out", save="Submit request")


@login_required
@require_POST
def exemption_action(request, pk, action):
    return decide(request, "exemptions", pk, action, "exemption request", exemption_to_attendance, "exemptions")


# ---------- personal travel allowance ----------
@login_required
def pta(request):
    rows = []
    for p in db().pta.find(scoped(request)).sort("created", -1).limit(500):
        p = clean(p)
        p["can_act"] = can_act(request, p)
        rows.append(p)
    return page(request, "pta.html", "pta", rows=rows)


@login_required
def pta_apply(request):
    form = PTAForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        c, me = form.cleaned_data, request.session["email"]
        month = c["month"].strip()
        if not MONTH_RE.fullmatch(month):
            form.add_error("month", "Choose a valid month.")
        elif month > month_str(date.today() + timedelta(days=31)):
            form.add_error("month", "That month is too far ahead.")
        elif db().pta.find_one({"email": me, "month": month, "status": {"$in": ["Pending", "Approved"]}}):
            form.add_error("month", "You already have a travel allowance request for that month.")
        if not form.errors:
            db().pta.insert_one(new_request(
                request, month=month, amount=float(c["amount"]), destination=c["destination"].strip(),
                purpose=c["purpose"].strip(), needs_hr=True))
            audit(request, "pta_requested", f"{month} {c['amount']}")
            messages.success(request, "Travel allowance request submitted.")
            return redirect("pta")
    return page(request, "form.html", "pta", status=form_feedback(request, form), form=form, back="pta",
                title="Apply for personal travel allowance", save="Submit request")


@login_required
@require_POST
def pta_action(request, pk, action):
    return decide(request, "pta", pk, action, "travel allowance request", None, "pta")


# ---------- payroll ----------
def build_payslip(emp, month):
    first, last = month_bounds(month)
    pta_total = sum(p["amount"] for p in db().pta.find(
        {"email": emp["email"], "month": month, "status": "Approved"}, {"amount": 1}))
    unpaid_days = 0
    for l in db().leaves.find({"email": emp["email"], "status": "Approved", "leave_type": "Unpaid",
                               "start_date": {"$lte": last.isoformat()}, "end_date": {"$gte": first.isoformat()}}):
        s = max(date.fromisoformat(l["start_date"]), first)
        e = min(date.fromisoformat(l["end_date"]), last)
        unpaid_days += sum(1 for _ in weekdays(s, e))
    basic = float(emp.get("salary", 0))                         # salary is treated as MONTHLY
    unpaid = round(basic / WORKING_DAYS_PER_MONTH * unpaid_days, 2)
    gross = round(basic - unpaid + pta_total, 2)
    pension = round((basic - unpaid) * PENSION_RATE, 2)
    tax = round(gross * TAX_RATE, 2)                            # replace with PAYE bands
    return {"month": month, "email": emp["email"], "name": f"{emp['first_name']} {emp['last_name']}",
            "work_id": emp.get("work_id", ""), "department": emp.get("department", ""),
            "position": emp.get("position", ""), "basic": basic, "pta": pta_total,
            "unpaid_days": unpaid_days, "unpaid_deduction": unpaid, "gross": gross,
            "pension": pension, "tax": tax, "net": round(gross - pension - tax, 2)}


@role_required(ADMINISTRATOR, ACCOUNTANT)
def payroll(request):
    month = request.GET.get("month", "")
    month = month if MONTH_RE.fullmatch(month) else month_str()
    rows = [clean(r) for r in db().payroll.find({"month": month}).sort("name", 1)]
    totals = {k: sum(r[k] for r in rows) for k in ("basic", "pta", "unpaid_deduction", "gross", "pension", "tax", "net")}
    final = bool(rows) and all(r["status"] == "Final" for r in rows)
    return page(request, "payroll.html", "payroll", rows=rows, month=month, totals=totals, final=final)


@role_required(ADMINISTRATOR, ACCOUNTANT)
@require_POST
def payroll_run(request):
    month = request.POST.get("month", "")
    if not MONTH_RE.fullmatch(month):
        messages.error(request, "Choose a valid month.")
        return redirect("payroll")
    if db().payroll.find_one({"month": month, "status": "Final"}):
        messages.error(request, f"Payroll for {month} is already finalised.")
        return redirect(f"{request.path.rsplit('run/', 1)[0]}?month={month}")
    now, n = datetime.now().isoformat(timespec="seconds"), 0
    for emp in db().employees.find({"status": "Active"}):
        db().payroll.update_one({"month": month, "email": emp["email"]},
                                {"$set": {**build_payslip(emp, month), "status": "Draft",
                                          "generated_by": request.session["name"], "generated_at": now}},
                                upsert=True)
        n += 1
    audit(request, "payroll_generated", f"{month} ({n} employees)")
    messages.success(request, f"Payroll for {month} generated for {n} employees.")
    return redirect(f"{request.path.rsplit('run/', 1)[0]}?month={month}")


@admin_required
@require_POST
def payroll_finalize(request):
    month = request.POST.get("month", "")
    if MONTH_RE.fullmatch(month):
        res = db().payroll.update_many({"month": month, "status": "Draft"},
                                       {"$set": {"status": "Final", "finalized_by": request.session["name"]}})
        audit(request, "payroll_finalized", f"{month} ({res.modified_count})")
        messages.success(request, f"Payroll for {month} is final. Payslips are now visible to employees.")
    return redirect(f"{request.path.rsplit('finalize/', 1)[0]}?month={month}")


@role_required(ADMINISTRATOR, ACCOUNTANT)
def payroll_export(request):
    month = request.GET.get("month", "")
    if not MONTH_RE.fullmatch(month):
        raise Http404
    out = [[r.get("work_id", ""), r["name"], r["department"], r["basic"], r["unpaid_days"], r["unpaid_deduction"],
            r["pta"], r["gross"], r["pension"], r["tax"], r["net"], r["status"]]
           for r in db().payroll.find({"month": month}).sort("name", 1)]
    audit(request, "payroll_export", month)
    return xlsx_response(f"payroll_{month}.xlsx",
                         ["Work ID", "Name", "Department", "Basic", "Unpaid days", "Unpaid deduction",
                          "Travel allowance", "Gross", "Pension", "Tax", "Net", "Status"], out)


@login_required
def payslips(request):
    rows = [clean(r) for r in db().payroll.find({"email": request.session["email"], "status": "Final"}).sort("month", -1)]
    return page(request, "payslips.html", "payslips", rows=rows)


@login_required
def payslip(request, pk):
    s = db().payroll.find_one({"_id": get_oid(pk)})
    staff = request.role in (ADMINISTRATOR, ACCOUNTANT)
    if not s or (not staff and (s["email"] != request.session["email"] or s["status"] != "Final")):
        raise Http404                                           # don't reveal other people's payslips
    audit(request, "payslip_viewed", f"{s['email']} {s['month']}")
    return page(request, "payslip.html", "payslips", s=clean(s))


@role_required(ADMINISTRATOR, ACCOUNTANT)
def invoice(request, month):
    if not MONTH_RE.fullmatch(month):
        raise Http404
    depts = {}
    for r in db().payroll.find({"month": month}):
        d = depts.setdefault(r.get("department") or "Unassigned",
                             {"name": r.get("department") or "Unassigned", "n": 0, "gross": 0, "pension": 0, "tax": 0, "net": 0})
        d["n"] += 1
        for k in ("gross", "pension", "tax", "net"):
            d[k] += r[k]
    if not depts:
        raise Http404
    lines = sorted(depts.values(), key=lambda d: d["name"])
    totals = {k: sum(d[k] for d in lines) for k in ("n", "gross", "pension", "tax", "net")}
    return page(request, "invoice.html", "payroll", month=month, lines=lines, totals=totals,
                number=f"INV-{month.replace('-', '')}", issued=date.today())


# ---------- recruitment ----------
def next_work_id():
    n = db().counters.find_one_and_update({"_id": "work_id"}, {"$inc": {"n": 1}}, upsert=True,
                                          return_document=ReturnDocument.AFTER)["n"]
    return f"FEW-{n:04d}"


@admin_required
def recruitment(request):
    stage = request.GET.get("stage", "")
    f = {"stage": stage} if stage in CANDIDATE_STAGES else {}
    rows = [clean(c) for c in db().candidates.find(f).sort("created", -1).limit(500)]
    return page(request, "recruitment.html", "recruitment", rows=rows, stage=stage, tabs=["", *CANDIDATE_STAGES])


@admin_required                                                  # STEP 1: create the candidate record
def candidate_add(request):
    form = CandidateForm(request.POST or None)
    form.fields["department"].choices = [(d["name"], d["name"]) for d in db().departments.find().sort("name", 1)]
    if request.method == "POST" and form.is_valid():
        c = {k: (v.strip() if isinstance(v, str) else v) for k, v in form.cleaned_data.items()}
        email = clean_email(c["email"])
        if not email:
            form.add_error("email", "Enter a valid email address.")
        elif db().candidates.find_one({"email": email, "stage": {"$in": ["Applied", "Interview"]}}):
            form.add_error("email", "This candidate is already in the pipeline.")
        elif db().employees.find_one({"email": email}):
            form.add_error("email", "An employee with this email already exists.")
        if not form.errors:
            c.update(email=email, expected_salary=float(c["expected_salary"]), stage="Applied", interviews=[],
                     created=datetime.now().isoformat(timespec="seconds"), created_by=request.session["name"])
            db().candidates.insert_one(c)
            audit(request, "candidate_created", email)
            messages.success(request, f"{c['first_name']} {c['last_name']} was added. Next: record an interview.")
            return redirect("recruitment")
    return page(request, "form.html", "recruitment", status=form_feedback(request, form), form=form,
                back="recruitment", title="New candidate", save="Create candidate")


@admin_required                                                  # STEP 2: record each interview
def interview_add(request, pk):
    cand = db().candidates.find_one({"_id": get_oid(pk)})
    if not cand:
        raise Http404
    if cand["stage"] not in ("Applied", "Interview"):
        messages.error(request, "This candidate is no longer in the interview process.")
        return redirect("recruitment")
    form = InterviewForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        c = form.cleaned_data
        entry = {"round": c["round"], "date": c["date"].isoformat(), "interviewer": c["interviewer"].strip(),
                 "outcome": c["outcome"], "score": c.get("score"), "notes": (c.get("notes") or "").strip(),
                 "recorded_by": request.session["name"],
                 "recorded_at": datetime.now().isoformat(timespec="seconds")}
        db().candidates.update_one({"_id": cand["_id"], "stage": {"$in": ["Applied", "Interview"]}},
                                   {"$push": {"interviews": entry}, "$set": {"stage": "Interview"}})
        audit(request, "interview_recorded", f"{cand['email']} {c['round']} {c['outcome']}")
        messages.success(request, f"{c['round']} interview recorded.")
        return redirect("recruitment")
    return page(request, "form.html", "recruitment", status=form_feedback(request, form), form=form,
                back="recruitment", title=f"Interview: {cand['first_name']} {cand['last_name']}", save="Save interview")


@admin_required                                                  # STEP 3: hire (creates employee + Work ID) or reject
@require_POST
def candidate_decide(request, pk, action):
    cand = db().candidates.find_one({"_id": get_oid(pk)})
    if not cand or action not in ("hire", "reject"):
        raise Http404
    now, who = datetime.now().isoformat(timespec="seconds"), request.session["name"]
    if cand["stage"] not in ("Applied", "Interview"):
        messages.info(request, "That candidate was already processed.")
    elif action == "reject":
        db().candidates.update_one({"_id": cand["_id"]}, {"$set": {"stage": "Rejected", "decided_by": who, "decided_at": now}})
        audit(request, "candidate_rejected", cand["email"])
        messages.success(request, "Candidate rejected.")
    elif cand["stage"] != "Interview" or not any(i["outcome"] == "Pass" for i in cand.get("interviews", [])):
        messages.error(request, "Record at least one passed interview before hiring.")
    else:
        work_id = next_work_id()
        claim = db().candidates.update_one({"_id": cand["_id"], "stage": "Interview"},
                                           {"$set": {"stage": "Hired", "work_id": work_id, "decided_by": who, "decided_at": now}})
        if claim.modified_count:
            data = {"first_name": cand["first_name"], "last_name": cand["last_name"], "email": cand["email"],
                    "phone": cand.get("phone", ""), "department": cand["department"], "position": cand["position"],
                    "hire_date": date.today().isoformat(), "salary": cand["expected_salary"],
                    "status": "Active", "role": EMPLOYEE, "work_id": work_id, "supervisor_email": ""}
            try:
                db().employees.insert_one(data)
            except DuplicateKeyError:
                db().candidates.update_one({"_id": cand["_id"]}, {"$set": {"stage": "Interview"}, "$unset": {"work_id": ""}})
                messages.error(request, "An employee with this email already exists.")
            else:
                audit(request, "candidate_hired", f"{cand['email']} {work_id}")
                _, mgr = department_manager(data["department"])
                mail.employee_welcome(data, mgr)
                messages.success(request, f"{cand['first_name']} {cand['last_name']} was hired. Work ID: {work_id}. "
                                          f"A welcome e-mail with registration instructions was sent.")
    return redirect("recruitment")




    # ---------- HR approvals (one page for every kind of pending request) ----------
HR_KINDS = {
    "leave": {"col": "leaves", "title": "Leave requests", "label": "leave request",
              "on_final": leave_to_attendance},
    "exemption": {"col": "exemptions", "title": "Missed clock-in / out", "label": "exemption request",
                  "on_final": exemption_to_attendance},
    "pta": {"col": "pta", "title": "Personal travel allowance", "label": "travel allowance request",
            "on_final": None},
}


@admin_required
def hr_approvals(request):
    """Every pending request from other employees, across all request types."""
    me = request.session["email"]
    flt = {"status": "Pending", "email": {"$ne": me}}           # HR can't review their own requests
    sections, total = [], 0
    for kind, cfg in HR_KINDS.items():
        rows = []
        for raw in db()[cfg["col"]].find(flt).sort("created", 1).limit(200):
            r = clean(raw)
            r["id"] = str(raw["_id"])                           # always available for the URL
            r["type"] = kind                                    # 'kind' is already used by exemptions
            rows.append(r)
        total += len(rows)
        sections.append({"kind": kind, "title": cfg["title"], "rows": rows})
    return page(request, "hr_approvals.html", "hr_approvals", sections=sections, total=total)

@admin_required
@require_POST
def hr_action(request, kind, pk, action):
    cfg = HR_KINDS.get(kind)
    if not cfg:
        raise Http404
    return decide(request, cfg["col"], pk, action, cfg["label"], cfg["on_final"], "hr_approvals")